"""
Dark-themed Personal Budget Tracker Excel Generator
Generates a styled .xlsx file compatible with Google Sheets import.
"""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

# ── Color Palette ──────────────────────────────────────────────────────────
BG         = "0D1117"
CARD_BG    = "161B22"
PURPLE     = "7C3AED"
GREEN      = "10B981"
AMBER      = "F59E0B"
RED        = "EF4444"
INPUT_BG   = "1C2333"
INPUT_TEXT  = "58A6FF"
GOLD       = "F0C040"
WHITE      = "FFFFFF"
BLACK      = "000000"
ALT_ROW    = "1A2332"
LIGHT_GRAY = "8B949E"

# ── Reusable style objects ─────────────────────────────────────────────────
fill_bg      = PatternFill("solid", fgColor=BG)
fill_card    = PatternFill("solid", fgColor=CARD_BG)
fill_input   = PatternFill("solid", fgColor=INPUT_BG)
fill_alt     = PatternFill("solid", fgColor=ALT_ROW)
fill_purple  = PatternFill("solid", fgColor=PURPLE)
fill_green   = PatternFill("solid", fgColor=GREEN)
fill_amber   = PatternFill("solid", fgColor=AMBER)
fill_red     = PatternFill("solid", fgColor=RED)
fill_gold    = PatternFill("solid", fgColor=GOLD)

font_white      = Font(color=WHITE, size=11)
font_white_bold = Font(color=WHITE, size=11, bold=True)
font_title      = Font(color=WHITE, size=16, bold=True)
font_input      = Font(color=INPUT_TEXT, size=11)
font_gold       = Font(color=GOLD, size=11, bold=True)
font_gold_big   = Font(color=GOLD, size=14, bold=True)
font_green      = Font(color=GREEN, size=11, bold=True)
font_red_bold   = Font(color=RED, size=11, bold=True)
font_amber      = Font(color=AMBER, size=11)
font_section    = Font(color=WHITE, size=12, bold=True)
font_formula    = Font(color=BLACK, size=11)

purple_side  = Side(style="thin", color=PURPLE)
border_input = Border(left=purple_side, right=purple_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left   = Alignment(horizontal="left", vertical="center")
align_right  = Alignment(horizontal="right", vertical="center")

USD     = '#,##0.00'
USD_NEG = '#,##0.00;[Red]-#,##0.00'
PCT     = '0.0%'
INT_FMT = '#,##0'


# ── Helpers ────────────────────────────────────────────────────────────────
def paint_bg(ws, max_row=80, max_col=20):
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill_bg
            cell.font = font_white


def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def section_header(ws, row, col_start, col_end, text, fill=fill_purple):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font_section
        cell.alignment = align_left
    ws.cell(row=row, column=col_start).value = text


def label_cell(ws, row, col, text, bold=False):
    cell = ws.cell(row=row, column=col)
    cell.value = text
    cell.font = font_white_bold if bold else font_white
    cell.alignment = align_left
    return cell


def input_cell(ws, row, col, value=None, fmt=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    cell.fill = fill_input
    cell.font = font_input
    cell.border = border_input
    cell.alignment = align_right
    if fmt:
        cell.number_format = fmt
    return cell


def formula_cell(ws, row, col, formula, fmt=None, font=None):
    cell = ws.cell(row=row, column=col)
    cell.value = formula
    cell.font = font or font_gold
    cell.fill = fill_card
    cell.alignment = align_right
    if fmt:
        cell.number_format = fmt
    return cell


def alt_row_fill(row_num):
    return fill_alt if row_num % 2 == 0 else fill_card


def style_row_bg(ws, row, col_start, col_end, fill):
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).fill = fill


# ═══════════════════════════════════════════════════════════════════════════
#  SHEET 1: Dashboard
# ═══════════════════════════════════════════════════════════════════════════
def build_dashboard(wb):
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_properties.tabColor = PURPLE
    ws.sheet_view.showGridLines = False
    paint_bg(ws, 55, 18)
    set_col_widths(ws, {
        'A': 3, 'B': 30, 'C': 18, 'D': 18, 'E': 18, 'F': 4,
        'G': 3, 'H': 28, 'I': 18, 'J': 18, 'K': 18, 'L': 18,
    })

    # Title
    ws.merge_cells('B1:E1')
    t = ws.cell(row=1, column=2)
    t.value = "Personal Budget Dashboard"
    t.font = font_title

    # ── Income Assumptions ──
    section_header(ws, 3, 2, 5, "INCOME ASSUMPTIONS")
    assumptions = [
        ("Annual Gross Salary",          78000,  USD),
        ("Pay Frequency (periods/yr)",   26,     INT_FMT),
        ("Savings per Paycheck",         700,    USD),
        ("Tulsa Remote Grant (annual)",  10000,  USD),
        ("OK State Tax Rate",            0.0475, PCT),
        ("Federal Effective Tax Rate",   0.18,   PCT),
    ]
    for i, (lbl, val, fmt) in enumerate(assumptions):
        r = 4 + i
        style_row_bg(ws, r, 2, 5, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)
        input_cell(ws, r, 3, val, fmt)
    # C4=salary C5=payfreq C6=savings/check C7=grant C8=state_tax C9=fed_tax

    # ── Income Summary ──
    section_header(ws, 11, 2, 5, "INCOME SUMMARY (MONTHLY)")
    summary = [
        ("Gross Monthly Income",         '=C4/12',       USD),
        ("Federal Tax",                  '=-(C4*C9)/12', USD),
        ("State Tax",                    '=-(C4*C8)/12', USD),
        ("Total Deductions",             '=C13+C14',     USD),
        ("Net Monthly Take-Home",        '=C12+C15',     USD),
        ("Tulsa Remote Grant (monthly)", '=C7/12',       USD),
        ("Net + Grant",                  '=C16+C17',     USD),
    ]
    for i, (lbl, fml, fmt) in enumerate(summary):
        r = 12 + i
        style_row_bg(ws, r, 2, 5, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)
        fnt = font_gold if r in (16, 18) else font_white
        formula_cell(ws, r, 3, fml, fmt, fnt)

    # ── Monthly Savings Allocation ──
    section_header(ws, 21, 2, 5, "MONTHLY SAVINGS ALLOCATION")
    for ci, hdr in enumerate(["Category", "Monthly", "Annual", "% of Income"], start=2):
        c = ws.cell(row=22, column=ci)
        c.value = hdr
        c.font = font_white_bold
        c.fill = fill_card
        c.alignment = align_center

    savings_items = [
        ("House Down Payment", 1200),
        ("Car Fund",           150),
        ("Travel Fund",        200),
        ("Emergency Fund",     300),
        ("General Savings",    100),
    ]
    for i, (lbl, amt) in enumerate(savings_items):
        r = 23 + i
        style_row_bg(ws, r, 2, 5, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)
        input_cell(ws, r, 3, amt, USD)
        formula_cell(ws, r, 4, f'=C{r}*12', USD, font_white)
        formula_cell(ws, r, 5, f'=C{r}/C16', PCT, font_white)

    # Total savings
    r_tot = 28
    style_row_bg(ws, r_tot, 2, 5, fill_card)
    label_cell(ws, r_tot, 2, "TOTAL SAVINGS", bold=True)
    formula_cell(ws, r_tot, 3, '=SUM(C23:C27)', USD, font_gold)
    formula_cell(ws, r_tot, 4, '=SUM(D23:D27)', USD, font_gold)
    formula_cell(ws, r_tot, 5, '=C28/C16', PCT, font_gold)

    # Remaining for living
    r_rem = 30
    style_row_bg(ws, r_rem, 2, 5, fill_card)
    label_cell(ws, r_rem, 2, "REMAINING FOR LIVING EXPENSES", bold=True)
    formula_cell(ws, r_rem, 3, '=C18-C28', USD, font_gold_big)

    # ── Savings Goals Tracker (right side) ──
    section_header(ws, 3, 8, 12, "SAVINGS GOALS TRACKER")
    goal_hdrs = ["Savings Bucket", "Target", "Saved So Far", "% Progress", "Months to Goal"]
    for ci, hdr in enumerate(goal_hdrs, start=8):
        c = ws.cell(row=4, column=ci)
        c.value = hdr
        c.font = font_white_bold
        c.fill = fill_card
        c.alignment = align_center

    goals = [
        ("House Down Payment", 60000, 5000),
        ("Car Fund",           15000, 2000),
        ("Travel Fund",        5000,  800),
        ("Emergency Fund",     10000, 3000),
        ("General Savings",    5000,  500),
    ]
    for i, (lbl, target, saved) in enumerate(goals):
        r = 5 + i
        style_row_bg(ws, r, 8, 12, alt_row_fill(r))
        label_cell(ws, r, 8, lbl)
        input_cell(ws, r, 9, target, USD)
        input_cell(ws, r, 10, saved, USD)
        formula_cell(ws, r, 11, f'=IF(I{r}=0,0,J{r}/I{r})', PCT, font_green)
        sav_row = 23 + i
        formula_cell(ws, r, 12,
                     f'=IF(C{sav_row}=0,"N/A",IF(J{r}>=I{r},"Done",(I{r}-J{r})/C{sav_row}))',
                     '0.0', font_amber)

    ws.freeze_panes = "B3"


# ═══════════════════════════════════════════════════════════════════════════
#  SHEET 2: Monthly Budget
# ═══════════════════════════════════════════════════════════════════════════
def build_monthly_budget(wb):
    ws = wb.create_sheet("Monthly Budget")
    ws.sheet_properties.tabColor = GREEN
    ws.sheet_view.showGridLines = False
    paint_bg(ws, 70, 12)
    set_col_widths(ws, {
        'A': 3, 'B': 30, 'C': 16, 'D': 16, 'E': 16, 'F': 14,
    })

    ws.merge_cells('B1:F1')
    ws.cell(row=1, column=2).value = "Monthly Budget"
    ws.cell(row=1, column=2).font = font_title

    label_cell(ws, 3, 2, "Month:", bold=True)
    input_cell(ws, 3, 3, "January")
    ws.cell(row=3, column=3).alignment = align_center

    for ci, hdr in enumerate(["Category", "Budgeted", "Actual", "Difference", "% Used"], start=2):
        c = ws.cell(row=5, column=ci)
        c.value = hdr
        c.font = font_white_bold
        c.fill = fill_card
        c.alignment = align_center

    sections = [
        ("HOUSING & UTILITIES", [
            ("Rent / Mortgage", 900), ("Electric", 80), ("Water / Trash", 40),
            ("Internet", 60), ("Renter's Insurance", 15),
        ]),
        ("FOOD", [
            ("Groceries", 300), ("Dining Out", 100), ("Coffee / Snacks", 30),
        ]),
        ("TRANSPORTATION", [
            ("Gas", 80), ("Car Insurance", 120), ("Car Maintenance", 50),
            ("Parking / Tolls", 15),
        ]),
        ("SUBSCRIPTIONS & TECH", [
            ("Phone Plan", 50), ("Streaming Services", 30),
            ("Software / Apps", 15), ("Cloud Storage", 5),
        ]),
        ("PERSONAL CARE", [
            ("Health Insurance (if any)", 0), ("Gym / Fitness", 35),
            ("Haircuts / Grooming", 25), ("Medical / Dental", 30),
        ]),
        ("ENTERTAINMENT & FUN", [
            ("Activities / Events", 50), ("Hobbies", 30), ("Gifts", 25),
        ]),
        ("TRAVEL", [
            ("Travel Fund Spending", 0),
        ]),
    ]

    row = 6
    first_data = None
    last_data = None
    for sec_name, items in sections:
        section_header(ws, row, 2, 6, sec_name, fill_purple)
        row += 1
        for lbl, budgeted in items:
            if first_data is None:
                first_data = row
            style_row_bg(ws, row, 2, 6, alt_row_fill(row))
            label_cell(ws, row, 2, lbl)
            input_cell(ws, row, 3, budgeted, USD)
            input_cell(ws, row, 4, 0, USD)
            formula_cell(ws, row, 5, f'=C{row}-D{row}', USD_NEG, font_white)
            formula_cell(ws, row, 6, f'=IF(C{row}=0,0,D{row}/C{row})', PCT, font_white)
            last_data = row
            row += 1
        row += 1

    # Totals
    section_header(ws, row, 2, 6, "TOTALS", fill_amber)
    row += 1
    style_row_bg(ws, row, 2, 6, fill_card)
    label_cell(ws, row, 2, "GRAND TOTAL", bold=True)
    formula_cell(ws, row, 3, f'=SUM(C{first_data}:C{last_data})', USD, font_gold)
    formula_cell(ws, row, 4, f'=SUM(D{first_data}:D{last_data})', USD, font_gold)
    formula_cell(ws, row, 5, f'=C{row}-D{row}', USD_NEG, font_gold)
    formula_cell(ws, row, 6, f'=IF(C{row}=0,0,D{row}/C{row})', PCT, font_gold)
    total_row = row

    row += 1
    style_row_bg(ws, row, 2, 6, fill_card)
    label_cell(ws, row, 2, "NET MONTHLY INCOME (from Dashboard)", bold=True)
    formula_cell(ws, row, 3, "='Dashboard'!C18", USD, font_green)

    row += 1
    style_row_bg(ws, row, 2, 6, fill_card)
    label_cell(ws, row, 2, "REMAINING AFTER EXPENSES", bold=True)
    formula_cell(ws, row, 3, f'=C{row-1}-D{total_row}', USD, font_gold_big)

    ws.freeze_panes = "B6"


# ═══════════════════════════════════════════════════════════════════════════
#  SHEET 3: Annual Plan
# ═══════════════════════════════════════════════════════════════════════════
def build_annual_plan(wb):
    ws = wb.create_sheet("Annual Plan")
    ws.sheet_properties.tabColor = AMBER
    ws.sheet_view.showGridLines = False
    paint_bg(ws, 30, 16)
    set_col_widths(ws, {'A': 3, 'B': 24})
    for i in range(3, 15):
        ws.column_dimensions[get_column_letter(i)].width = 14

    ws.merge_cells('B1:N1')
    ws.cell(row=1, column=2).value = "12-Month Savings Projection"
    ws.cell(row=1, column=2).font = font_title

    months = ["Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
              "Jan", "Feb", "Mar", "Apr", "May", "Jun"]

    section_header(ws, 3, 2, 14, "MONTHLY SAVINGS PROJECTION")
    ws.cell(row=4, column=2).value = "Savings Bucket"
    ws.cell(row=4, column=2).font = font_white_bold
    ws.cell(row=4, column=2).fill = fill_card
    for i, m in enumerate(months):
        c = ws.cell(row=4, column=3 + i)
        c.value = m
        c.font = font_white_bold
        c.fill = fill_card
        c.alignment = align_center

    buckets = [
        ("House Down Payment", 23),
        ("Car Fund",           24),
        ("Travel Fund",        25),
        ("Emergency Fund",     26),
        ("General Savings",    27),
    ]
    for i, (lbl, dash_row) in enumerate(buckets):
        r = 5 + i
        style_row_bg(ws, r, 2, 14, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)
        for m in range(12):
            formula_cell(ws, r, 3 + m, f"='Dashboard'!C{dash_row}", USD, font_white)

    # Monthly total
    r_mt = 10
    style_row_bg(ws, r_mt, 2, 14, fill_card)
    label_cell(ws, r_mt, 2, "Monthly Total", bold=True)
    for m in range(12):
        cl = get_column_letter(3 + m)
        formula_cell(ws, r_mt, 3 + m, f'=SUM({cl}5:{cl}9)', USD, font_gold)

    # Cumulative total
    r_ct = 11
    style_row_bg(ws, r_ct, 2, 14, fill_card)
    label_cell(ws, r_ct, 2, "Cumulative Total", bold=True)
    formula_cell(ws, r_ct, 3, '=C10', USD, font_gold)
    for m in range(1, 12):
        col = 3 + m
        prev = get_column_letter(col - 1)
        cur = get_column_letter(col)
        formula_cell(ws, r_ct, col, f'={prev}{r_ct}+{cur}{r_mt}', USD, font_gold)

    # Tulsa Remote Grant
    section_header(ws, 13, 2, 14, "TULSA REMOTE GRANT ALLOCATION")
    r_gr = 14
    style_row_bg(ws, r_gr, 2, 14, fill_card)
    label_cell(ws, r_gr, 2, "Grant Applied to Rent", bold=True)
    for m in range(12):
        formula_cell(ws, r_gr, 3 + m, "='Dashboard'!C7/12", USD, font_green)

    ws.freeze_panes = "C5"


# ═══════════════════════════════════════════════════════════════════════════
#  SHEET 4: City Comparison
# ═══════════════════════════════════════════════════════════════════════════
def build_city_comparison(wb):
    ws = wb.create_sheet("City Comparison")
    ws.sheet_properties.tabColor = RED
    ws.sheet_view.showGridLines = False
    paint_bg(ws, 30, 10)
    set_col_widths(ws, {
        'A': 3, 'B': 24, 'C': 16, 'D': 16, 'E': 18, 'F': 20,
    })

    ws.merge_cells('B1:F1')
    ws.cell(row=1, column=2).value = "City Cost-of-Living Comparison"
    ws.cell(row=1, column=2).font = font_title

    section_header(ws, 3, 2, 6, "COST COMPARISON")
    for ci, hdr in enumerate(["Category", "Dallas", "Tulsa", "NW Arkansas", "Best Option"], start=2):
        c = ws.cell(row=4, column=ci)
        c.value = hdr
        c.font = font_white_bold
        c.fill = fill_card
        c.alignment = align_center

    categories = [
        # (label, dallas, tulsa, nw_ark, is_formula, fmt)
        ("Rent (1BR)",        1200,   900,    950,   False, USD),
        ("Grant Offset",      0,      -833,   0,     False, USD),
        ("Effective Rent",    None,   None,   None,  True,  USD),
        ("Utilities",         150,    120,    125,   False, USD),
        ("Groceries",         350,    280,    290,   False, USD),
        ("Gas (monthly)",     100,    75,     80,    False, USD),
        ("Car Insurance",     140,    100,    110,   False, USD),
        ("Dining Out",        120,    80,     85,    False, USD),
        ("State Tax Rate",    0,      0.0475, 0.047, False, PCT),
        ("Avg Home Price",    350000, 210000, 240000,False, '#,##0'),
    ]

    for i, (lbl, dal, tul, nwa, is_fml, fmt) in enumerate(categories):
        r = 5 + i
        style_row_bg(ws, r, 2, 6, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)

        if is_fml:
            # Effective Rent = Rent + Grant Offset (2 rows above + 1 row above)
            formula_cell(ws, r, 3, f'=C{r-2}+C{r-1}', fmt, font_white)
            formula_cell(ws, r, 4, f'=D{r-2}+D{r-1}', fmt, font_white)
            formula_cell(ws, r, 5, f'=E{r-2}+E{r-1}', fmt, font_white)
        else:
            input_cell(ws, r, 3, dal, fmt)
            input_cell(ws, r, 4, tul, fmt)
            input_cell(ws, r, 5, nwa, fmt)

        # Best Option column
        formula_cell(ws, r, 6,
                     f'=IF(AND(C{r}<=D{r},C{r}<=E{r}),"Dallas",'
                     f'IF(AND(D{r}<=C{r},D{r}<=E{r}),"Tulsa","NW Arkansas"))',
                     None, font_green)

    # Conditional formatting: highlight lowest value green bold in each row
    for i in range(len(categories)):
        r = 5 + i
        for col in range(3, 6):
            cl = get_column_letter(col)
            others = [get_column_letter(c) for c in range(3, 6) if c != col]
            ws.conditional_formatting.add(
                f'{cl}{r}',
                FormulaRule(
                    formula=[f'=AND({cl}{r}<={others[0]}{r},{cl}{r}<={others[1]}{r})'],
                    font=Font(color=GREEN, bold=True),
                )
            )

    ws.freeze_panes = "B5"


# ═══════════════════════════════════════════════════════════════════════════
#  SHEET 5: Net Worth Tracker
# ═══════════════════════════════════════════════════════════════════════════
def build_net_worth(wb):
    ws = wb.create_sheet("Net Worth Tracker")
    ws.sheet_properties.tabColor = GOLD
    ws.sheet_view.showGridLines = False
    paint_bg(ws, 40, 10)
    set_col_widths(ws, {'A': 3, 'B': 30, 'C': 20})

    ws.merge_cells('B1:C1')
    ws.cell(row=1, column=2).value = "Net Worth Tracker"
    ws.cell(row=1, column=2).font = font_title

    # Assets
    section_header(ws, 3, 2, 3, "ASSETS", fill_green)
    assets = [
        ("Checking Account",        2500),
        ("Savings / HYSA",          5000),
        ("House Down Payment Fund", 5000),
        ("Car Fund",                2000),
        ("Travel Fund",             800),
        ("Emergency Fund",          3000),
        ("General Savings",         500),
        ("401(k)",                  8000),
        ("Investments / Brokerage", 1500),
    ]
    for i, (lbl, val) in enumerate(assets):
        r = 4 + i
        style_row_bg(ws, r, 2, 3, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)
        input_cell(ws, r, 3, val, USD)

    r_at = 4 + len(assets)  # 13
    style_row_bg(ws, r_at, 2, 3, fill_card)
    label_cell(ws, r_at, 2, "TOTAL ASSETS", bold=True)
    formula_cell(ws, r_at, 3, f'=SUM(C4:C{r_at - 1})', USD, font_green)

    # Liabilities
    r_lh = r_at + 2  # 15
    section_header(ws, r_lh, 2, 3, "LIABILITIES", fill_red)
    liabilities = [
        ("Car Loan",      0),
        ("Student Loans",  0),
        ("Credit Cards",   0),
        ("Other Debt",     0),
    ]
    for i, (lbl, val) in enumerate(liabilities):
        r = r_lh + 1 + i
        style_row_bg(ws, r, 2, 3, alt_row_fill(r))
        label_cell(ws, r, 2, lbl)
        input_cell(ws, r, 3, val, USD)

    r_lf = r_lh + 1
    r_ll = r_lh + len(liabilities)
    r_lt = r_ll + 1
    style_row_bg(ws, r_lt, 2, 3, fill_card)
    label_cell(ws, r_lt, 2, "TOTAL LIABILITIES", bold=True)
    formula_cell(ws, r_lt, 3, f'=SUM(C{r_lf}:C{r_ll})', USD, font_red_bold)

    # Net Worth
    r_nwh = r_lt + 2
    section_header(ws, r_nwh, 2, 3, "NET WORTH", fill_purple)
    r_nw = r_nwh + 1
    style_row_bg(ws, r_nw, 2, 3, fill_card)
    label_cell(ws, r_nw, 2, "Total Net Worth", bold=True)
    formula_cell(ws, r_nw, 3, f'=C{r_at}-C{r_lt}', USD, font_gold_big)
    ws.row_dimensions[r_nw].height = 30

    ws.freeze_panes = "B3"


# ═══════════════════════════════════════════════════════════════════════════
#  Main
# ═══════════════════════════════════════════════════════════════════════════
def main():
    wb = openpyxl.Workbook()
    build_dashboard(wb)
    build_monthly_budget(wb)
    build_annual_plan(wb)
    build_city_comparison(wb)
    build_net_worth(wb)

    output = "budget_tracker.xlsx"
    wb.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
