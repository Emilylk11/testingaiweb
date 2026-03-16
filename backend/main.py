"""FastAPI backend for the Budget Tracker web app."""

import json
import io
from pathlib import Path

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import openpyxl

app = FastAPI(title="Budget Tracker API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

DATA_FILE = Path(__file__).parent / "data.json"


def load_data() -> dict:
    return json.loads(DATA_FILE.read_text())


def save_data(data: dict):
    DATA_FILE.write_text(json.dumps(data, indent=2))


# ── GET endpoints ──────────────────────────────────────────────────────────

@app.get("/api/data")
def get_all_data():
    """Return the full budget dataset with computed summaries."""
    data = load_data()
    inc = data["income"]

    gross_monthly = inc["annual_salary"] / 12
    fed_tax = inc["annual_salary"] * inc["federal_tax_rate"] / 12
    state_tax = inc["annual_salary"] * inc["state_tax_rate"] / 12
    net_monthly = gross_monthly - fed_tax - state_tax
    grant_monthly = inc["tulsa_grant_annual"] / 12
    net_plus_grant = net_monthly + grant_monthly

    total_savings = sum(s["monthly"] for s in data["savings_allocation"])
    remaining = net_plus_grant - total_savings

    data["computed"] = {
        "gross_monthly": round(gross_monthly, 2),
        "federal_tax": round(fed_tax, 2),
        "state_tax": round(state_tax, 2),
        "net_monthly": round(net_monthly, 2),
        "grant_monthly": round(grant_monthly, 2),
        "net_plus_grant": round(net_plus_grant, 2),
        "total_savings_monthly": round(total_savings, 2),
        "remaining_for_living": round(remaining, 2),
    }

    # Budget totals
    total_budgeted = 0
    total_actual = 0
    for cat in data["monthly_budget"]["categories"]:
        for item in cat["items"]:
            total_budgeted += item["budgeted"]
            total_actual += item["actual"]
    data["computed"]["total_budgeted"] = round(total_budgeted, 2)
    data["computed"]["total_actual"] = round(total_actual, 2)
    data["computed"]["budget_remaining"] = round(net_plus_grant - total_actual, 2)

    # Net worth
    total_assets = sum(a["value"] for a in data["net_worth"]["assets"])
    total_liabilities = sum(l["value"] for l in data["net_worth"]["liabilities"])
    data["computed"]["total_assets"] = round(total_assets, 2)
    data["computed"]["total_liabilities"] = round(total_liabilities, 2)
    data["computed"]["net_worth"] = round(total_assets - total_liabilities, 2)

    return data


# ── PUT endpoints ──────────────────────────────────────────────────────────

class IncomeUpdate(BaseModel):
    annual_salary: float
    pay_frequency: int
    savings_per_paycheck: float
    tulsa_grant_annual: float
    state_tax_rate: float
    federal_tax_rate: float


@app.put("/api/income")
def update_income(body: IncomeUpdate):
    data = load_data()
    data["income"] = body.model_dump()
    save_data(data)
    return {"status": "ok"}


class SavingsItem(BaseModel):
    category: str
    monthly: float


@app.put("/api/savings-allocation")
def update_savings_allocation(body: list[SavingsItem]):
    data = load_data()
    data["savings_allocation"] = [s.model_dump() for s in body]
    save_data(data)
    return {"status": "ok"}


class GoalItem(BaseModel):
    category: str
    target: float
    saved: float


@app.put("/api/savings-goals")
def update_savings_goals(body: list[GoalItem]):
    data = load_data()
    data["savings_goals"] = [g.model_dump() for g in body]
    save_data(data)
    return {"status": "ok"}


class BudgetItemUpdate(BaseModel):
    name: str
    budgeted: float
    actual: float


class BudgetCategoryUpdate(BaseModel):
    section: str
    items: list[BudgetItemUpdate]


class MonthlyBudgetUpdate(BaseModel):
    month: str
    categories: list[BudgetCategoryUpdate]


@app.put("/api/monthly-budget")
def update_monthly_budget(body: MonthlyBudgetUpdate):
    data = load_data()
    data["monthly_budget"] = body.model_dump()
    save_data(data)
    return {"status": "ok"}


class AssetItem(BaseModel):
    name: str
    value: float


class NetWorthUpdate(BaseModel):
    assets: list[AssetItem]
    liabilities: list[AssetItem]


@app.put("/api/net-worth")
def update_net_worth(body: NetWorthUpdate):
    data = load_data()
    data["net_worth"] = body.model_dump()
    save_data(data)
    return {"status": "ok"}


class CityRow(BaseModel):
    category: str
    dallas: float
    tulsa: float
    nw_arkansas: float


@app.put("/api/city-comparison")
def update_city_comparison(body: list[CityRow]):
    data = load_data()
    data["city_comparison"] = [c.model_dump() for c in body]
    save_data(data)
    return {"status": "ok"}


# ── Excel export ───────────────────────────────────────────────────────────

@app.get("/api/export")
def export_excel():
    """Generate and return the budget_tracker.xlsx."""
    import subprocess, sys
    script = Path(__file__).parent.parent / "budget_tracker.py"
    out = Path(__file__).parent.parent / "budget_tracker.xlsx"
    subprocess.run([sys.executable, str(script)], check=True)
    buf = io.BytesIO(out.read_bytes())
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=budget_tracker.xlsx"},
    )


# ── Excel import ───────────────────────────────────────────────────────────

@app.post("/api/import")
async def import_excel(file: UploadFile = File(...)):
    """Parse an uploaded .xlsx and update data.json with extracted values."""
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    data = load_data()

    if "Dashboard" in wb.sheetnames:
        ws = wb["Dashboard"]
        # Income assumptions: C4-C9
        mapping = [
            ("annual_salary", 4), ("pay_frequency", 5),
            ("savings_per_paycheck", 6), ("tulsa_grant_annual", 7),
            ("state_tax_rate", 8), ("federal_tax_rate", 9),
        ]
        for key, row in mapping:
            val = ws.cell(row=row, column=3).value
            if val is not None:
                data["income"][key] = float(val) if not isinstance(val, (int, float)) else val

        # Savings allocation: C23-C27
        for i in range(5):
            val = ws.cell(row=23 + i, column=3).value
            if val is not None and i < len(data["savings_allocation"]):
                data["savings_allocation"][i]["monthly"] = float(val)

        # Goals: I5-J9
        for i in range(5):
            target = ws.cell(row=5 + i, column=9).value
            saved = ws.cell(row=5 + i, column=10).value
            if i < len(data["savings_goals"]):
                if target is not None:
                    data["savings_goals"][i]["target"] = float(target)
                if saved is not None:
                    data["savings_goals"][i]["saved"] = float(saved)

    if "Net Worth Tracker" in wb.sheetnames:
        ws = wb["Net Worth Tracker"]
        for i in range(len(data["net_worth"]["assets"])):
            val = ws.cell(row=4 + i, column=3).value
            if val is not None:
                data["net_worth"]["assets"][i]["value"] = float(val)
        for i in range(len(data["net_worth"]["liabilities"])):
            val = ws.cell(row=16 + i, column=3).value
            if val is not None:
                data["net_worth"]["liabilities"][i]["value"] = float(val)

    save_data(data)
    return {"status": "ok", "message": "Data imported successfully"}
